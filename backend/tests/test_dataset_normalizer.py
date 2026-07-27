from pathlib import Path

from openpyxl import Workbook
from PIL import Image

from app.ingestion.delimited import (
    extract_delimited_table,
)
from app.ingestion.detector import (
    inspect_file,
)
from app.ingestion.normalizer import (
    apply_human_review,
    build_unified_datasets,
)
from app.ingestion.pipeline import (
    run_ingestion_pipeline,
)
from app.ingestion.schemas import (
    HumanReviewDecision,
    HumanReviewRequest,
    DatasetCorrection,
    ReviewStatus,
    SourceKind,
)


def test_excel_pipeline_creates_dataset(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "sales.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sales"

    worksheet.append([
        "月份",
        "營收",
    ])

    worksheet.append([
        "2026-01",
        100000,
    ])

    worksheet.append([
        "2026-02",
        120000,
    ])

    worksheet.append([
        "2026-03",
        135000,
    ])

    workbook.save(file_path)
    workbook.close()

    first_result = run_ingestion_pipeline(
        file_path
    )

    second_result = run_ingestion_pipeline(
        file_path
    )

    assert len(first_result.datasets) == 1

    dataset = first_result.datasets[0]

    assert (
        dataset.source_kind
        == SourceKind.EXCEL
    )

    assert dataset.row_count == 3

    assert (
        dataset.requires_human_review
        is False
    )

    assert (
        dataset.review_status
        == ReviewStatus.NOT_REQUIRED
    )

    # 同一份資料必須產生穩定 ID。
    assert (
        dataset.dataset_id
        == second_result.datasets[0].dataset_id
    )


def test_image_source_requires_review(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "source.csv"

    csv_path.write_text(
        "月份,營收\n"
        "2026-01,100000\n"
        "2026-02,120000\n"
        "2026-03,135000\n",
        encoding="utf-8-sig",
    )

    extraction = extract_delimited_table(
        csv_path
    )

    table = extraction.tables[0]

    visual_table = table.model_copy(
        update={
            "filename": "table.png",
            "sheet_name":
                "visual_page_1_table_1",
        }
    )

    extraction = extraction.model_copy(
        update={
            "filename": "table.png",
            "tables": [visual_table],
        }
    )

    image_path = tmp_path / "table.png"

    image = Image.new(
        "RGB",
        (100, 100),
        "white",
    )

    image.save(image_path)

    inspection = inspect_file(
        image_path
    )

    datasets = build_unified_datasets(
        extraction=extraction,
        inspection=inspection,
    )

    assert len(datasets) == 1

    dataset = datasets[0]

    assert (
        dataset.source_kind
        == SourceKind.OCR_IMAGE_TABLE
    )

    assert (
        dataset.requires_human_review
        is True
    )

    assert (
        dataset.review_status
        == ReviewStatus.PENDING
    )


def test_apply_human_review(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "source.csv"

    csv_path.write_text(
        "月份,營收\n"
        "2026-01,100000\n"
        "2026-02,120000\n"
        "2026-03,135000\n",
        encoding="utf-8-sig",
    )

    extraction = extract_delimited_table(
        csv_path
    )

    image_path = tmp_path / "table.png"

    image = Image.new(
        "RGB",
        (100, 100),
        "white",
    )

    image.save(image_path)

    inspection = inspect_file(
        image_path
    )

    table = extraction.tables[0].model_copy(
        update={
            "filename": "table.png",
            "sheet_name":
                "visual_page_1_table_1",
        }
    )

    extraction = extraction.model_copy(
        update={
            "filename": "table.png",
            "tables": [table],
        }
    )

    dataset = build_unified_datasets(
        extraction=extraction,
        inspection=inspection,
    )[0]

    reviewed = apply_human_review(
        dataset=dataset,
        review=HumanReviewRequest(
            decision=(
                HumanReviewDecision.APPROVE
            ),
            reviewer="Brian",
            notes="已核對原始圖片",
            corrections=[
                DatasetCorrection(
                    record_index=0,
                    column_key="營收",
                    corrected_value=110000,
                    note="OCR 原本辨識錯誤",
                )
            ],
        ),
    )

    assert (
        reviewed.review_status
        == ReviewStatus.APPROVED
    )

    assert (
        reviewed.requires_human_review
        is False
    )

    assert (
        reviewed.records[0]
        .values["營收"]
        .value
        == 110000
    )

    assert (
        reviewed.records[0]
        .values["營收"]
        .confidence
        == 1.0
    )